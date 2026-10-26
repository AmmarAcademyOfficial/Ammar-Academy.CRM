"""
CRM changes, Part C browser test (26 Oct 2026): Download Excel on the Leads and Enrollments lists —
exactly what the filters show, real dates and numbers, three sheets for enrollments, no phone numbers
for the accountant, who may download, and the audit log line for each download. Synthetic data only.

Runs against the production build and a freshly seeded local database:
  python3 tests/e2e/excel_downloads_e2e.py [screenshot_dir]
"""
import datetime as dt
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl
from playwright.sync_api import expect, sync_playwright
import hydration  # noqa: F401  (waits for the page to be interactive after each load)

BASE = "http://localhost:3000"
ROOT = Path(__file__).parents[2]
SHOTS = Path(sys.argv[1]) if len(sys.argv) > 1 else None
results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))
    if not cond:
        raise SystemExit(1)


def shot(page, name, full=False):
    if SHOTS:
        SHOTS.mkdir(parents=True, exist_ok=True)
        page.wait_for_timeout(400)
        page.screenshot(path=str(SHOTS / f"{name}.png"), full_page=full)


def sql(query, params=None):
    js = (
        "const pg=require('pg');const c=new pg.Client('postgres://postgres:postgres@localhost:54329/crm');"
        "c.connect().then(()=>c.query(process.argv[1], JSON.parse(process.argv[2]))).then(r=>{process.stdout.write(JSON.stringify(r.rows));return c.end()})"
        ".catch(e=>{console.error(e.message);process.exit(1)})"
    )
    out = subprocess.run(["node", "-e", js, query, json.dumps(params or [])], capture_output=True, text=True, cwd=ROOT)
    if out.returncode:
        raise RuntimeError(out.stderr)
    return json.loads(out.stdout or "[]")


def sign_in(ctx, name, path):
    page = ctx.new_page()
    page.goto(f"{BASE}/login")
    page.get_by_role("button", name=name).click()
    page.wait_for_url(f"**{path}*")
    return page


def download(page):
    with page.expect_download() as dl:
        page.get_by_role("button", name="Download Excel").click()
    expect(page.get_by_text("Excel file downloaded").last).to_be_visible()
    path = Path(tempfile.mkdtemp()) / dl.value.suggested_filename  # openpyxl needs the .xlsx name
    dl.value.save_as(path)
    return dl.value.suggested_filename, openpyxl.load_workbook(path)


def data_rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and not str(r[0]).startswith("Total ")]


def last_audit():
    return sql("select a.target, p.full_name from public.audit_log a join public.profiles p on p.id = a.actor_id where a.action = 'Downloaded Excel' order by a.created_at desc limit 1")[0]


ALL_LEADS = sql("select count(*)::int as n from public.inquiries")[0]["n"]
ALL_ENROLLMENTS = sql("select count(*)::int as n from public.enrollments")[0]["n"]
ALL_PAYMENTS = sql("select count(*)::int as n from public.payments")[0]["n"]
ALL_DUES = sql("select count(*)::int as n from public.payment_dues")[0]["n"]

with sync_playwright() as p:
    browser = p.chromium.launch()

    # ── 1. Manager: Leads, all and filtered ─────────────────────────────────
    mctx = browser.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
    m = sign_in(mctx, "Hana Farouk", "/overview")
    m.goto(f"{BASE}/leads")
    expect(m.get_by_role("button", name="Download Excel")).to_be_enabled()
    shot(m, "01-leads-button")
    name, wb = download(m)
    ws = wb["Leads"]
    headers = [c.value for c in ws[1]]
    rows = data_rows(ws)
    check("A manager downloads every lead as a real Excel file", name.startswith("leads-") and name.endswith(".xlsx") and len(rows) == ALL_LEADS,
          f"{name}, {len(rows)} of {ALL_LEADS}")
    check("The Leads sheet has the planned columns", headers == ["Name", "Phone", "Student", "Program", "Source", "Stage", "Last outcome", "Sales agent", "Tags",
                                                               "Created", "Assigned", "First contact", "Last contact", "Next follow-up", "Course of interest", "Enrolled"], json.dumps(headers))
    created = [r[9] for r in rows]
    check("Dates are real Excel dates and phones are included", all(isinstance(x, dt.datetime) for x in created) and all(str(r[1]).startswith("+") for r in rows))
    check("The header row is frozen with filter arrows", ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:P"))
    a = last_audit()
    check("The download is in the audit log", a == {"target": f"Leads · {ALL_LEADS} rows · no filters", "full_name": "Hana Farouk"}, json.dumps(a))

    m.get_by_label("Search by name or phone").fill("Sara")
    expect(m.get_by_text("matching inquir")).to_be_visible()
    expect(m.get_by_text(f"{ALL_LEADS} matching inquir")).to_have_count(0)  # the search has applied
    shown = int(m.get_by_text("matching inquir").inner_text().split()[0])
    _, wb = download(m)
    rows = data_rows(wb["Leads"])
    check("A filtered download holds exactly what the list shows", len(rows) == shown and all("sara" in (str(r[0]) + str(r[2])).lower() for r in rows), f"{len(rows)} vs {shown}")
    check("The audit line names the filters", last_audit()["target"] == f'Leads · {shown} row{"" if shown == 1 else "s"} · search "Sara"', last_audit()["target"])

    # ── 2. Manager: Enrollments with Payments and Payment plan ──────────────
    m.goto(f"{BASE}/customers")
    expect(m.get_by_role("button", name="Download Excel")).to_be_enabled()
    check("The old CSV button is replaced by Download Excel", m.get_by_role("button", name="Download CSV").count() == 0)
    name, wb = download(m)
    shot(m, "02-customers-downloaded")
    check("The enrollments file has three sheets", wb.sheetnames == ["Enrollments", "Payments", "Payment plan"], json.dumps(wb.sheetnames))
    enr, pay, plan = data_rows(wb["Enrollments"]), data_rows(wb["Payments"]), data_rows(wb["Payment plan"])
    check("Every enrollment, payment and planned payment is in the file", (len(enr), len(pay), len(plan)) == (ALL_ENROLLMENTS, ALL_PAYMENTS, ALL_DUES),
          f"{len(enr)}/{ALL_ENROLLMENTS}, {len(pay)}/{ALL_PAYMENTS}, {len(plan)}/{ALL_DUES}")
    eh = [c.value for c in wb["Enrollments"][1]]
    paid_col, cur_col = eh.index("Paid"), eh.index("Currency")
    db_paid = {r["currency"]: float(r["paid"]) for r in sql("""select e.currency, coalesce(sum(p.amount) filter (where p.status = 'confirmed'), 0) as paid
                                                                 from public.enrollments e left join public.payments p on p.enrollment_id = e.id group by e.currency""")}
    file_paid = {}
    for r in enr:
        file_paid[r[cur_col]] = file_paid.get(r[cur_col], 0) + float(r[paid_col] or 0)
    check("Amounts are numbers and add up to the database totals", all(isinstance(r[paid_col], (int, float)) for r in enr)
          and all(abs(file_paid.get(c, 0) - v) < 0.01 for c, v in db_paid.items()), json.dumps([file_paid, db_paid]))
    totals = [r for r in wb["Enrollments"].iter_rows(values_only=True) if r[0] and str(r[0]).startswith("Total ")]
    check("A totals row per currency sits under the money columns", len(totals) == len(db_paid), json.dumps([t[0] for t in totals]))
    ph = [c.value for c in wb["Payments"][1]]
    check("Payments have the planned columns", ph[:8] == ["Customer", "Phone", "Student", "Course", "Date paid", "Amount", "Currency", "Type"] and "Confirmed or rejected by" in ph, json.dumps(ph))
    check("The download is logged with the row count", last_audit()["target"].startswith(f"Enrollments · {ALL_ENROLLMENTS} rows"), last_audit()["target"])

    # ── 3. Accountant: no phone numbers anywhere ────────────────────────────
    actx = browser.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
    a = sign_in(actx, "Karim Fawzy", "/accounting")
    a.goto(f"{BASE}/accounting?tab=customers")
    _, wb = download(a)
    text = " ".join(str(c.value) for sh in wb.worksheets for row in sh.iter_rows() for c in row if c.value is not None)
    check("The accountant's file has no phone numbers", all("Phone" not in [c.value for c in sh[1]] for sh in wb.worksheets) and not re.search(r"\+\d{8,}", text))
    check("The accountant's download is logged too", last_audit()["full_name"] == "Karim Fawzy")

    # ── 4. Support and sales cannot download ────────────────────────────────
    uctx = browser.new_context(viewport={"width": 1440, "height": 900})
    u = sign_in(uctx, "Laila Ibrahim", "/my-work")
    u.goto(f"{BASE}/customers")
    expect(u.get_by_role("button", name="Filters")).to_be_visible()
    body = {"filters": {}, "sort": {"key": "enrolled", "dir": "desc"}}
    check("Support has no download button and the server refuses them", u.get_by_role("button", name="Download Excel").count() == 0
          and u.request.post(f"{BASE}/api/enrollments/export", data=body).status == 403)
    sctx = browser.new_context(viewport={"width": 1440, "height": 900})
    s = sign_in(sctx, "Nour Hassan", "/my-work")
    s.goto(f"{BASE}/leads")
    expect(s.get_by_role("button", name="Filters")).to_be_visible()
    lead_body = {"filters": {}, "sort": {"key": "created", "dir": "desc"}}
    check("Sales agents have no download button and the server refuses them", s.get_by_role("button", name="Download Excel").count() == 0
          and s.request.post(f"{BASE}/api/leads/export", data=lead_body).status == 403
          and a.request.post(f"{BASE}/api/leads/export", data=lead_body).status == 403)

    # ── 5. The download history in Settings → Audit ─────────────────────────
    m.goto(f"{BASE}/settings/audit")
    expect(m.get_by_text("Downloaded Excel").first).to_be_visible()
    check("Managers see the download history in the audit log", m.get_by_text("Downloaded Excel").count() >= 4)
    shot(m, "03-audit")
    browser.close()

print(f"\n{sum(ok for _, ok in results)}/{len(results)} checks passed")
